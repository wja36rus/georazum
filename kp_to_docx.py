#!/usr/bin/env python3
"""Конвертер kp-slides.html -> kp.docx (A4, редактируемый Word-документ)."""

import re
from html.parser import HTMLParser

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Mm, Pt, RGBColor

GREEN = RGBColor(0x0F, 0x8A, 0x5F)
INK = RGBColor(0x0E, 0x1A, 0x13)
MUTED = RGBColor(0x5B, 0x6B, 0x60)
AMBER = RGBColor(0xC9, 0x79, 0x1F)

VOID = {"br", "hr", "img", "meta", "link", "input"}


class Node:
    def __init__(self, tag=None, attrs=None):
        self.tag = tag
        self.attrs = dict(attrs or {})
        self.children = []

    @property
    def cls(self):
        return set(self.attrs.get("class", "").split())

    def kids(self):
        return [c for c in self.children if isinstance(c, Node)]

    def text(self):
        s = "".join(c if isinstance(c, str) else c.text() for c in self.children)
        s = re.sub(r"[ \t]+", " ", s)
        return re.sub(r" *\n *", "\n", s).strip()


class Tree(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node("root")
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        if tag in VOID:
            if tag == "br":
                self.stack[-1].children.append("\n")
            return
        node = Node(tag, attrs)
        self.stack[-1].children.append(node)
        self.stack.append(node)

    def handle_endtag(self, tag):
        if tag in VOID:
            return
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                return

    def handle_data(self, data):
        self.stack[-1].children.append(data)


def add_runs(p, node, bold=False, italic=False, size=None, color=None):
    pieces = []

    def walk(n, b, i):
        for c in n.children:
            if isinstance(c, str):
                t = re.sub(r"\s+", " ", c)
                if t:
                    pieces.append((t, b, i))
            elif c.tag == "br":
                pieces.append((" ", b, i))
            elif c.tag in ("b", "strong"):
                walk(c, True, i)
            elif c.tag in ("em", "i"):
                walk(c, b, True)
            else:
                walk(c, b, i)

    walk(node, bold, italic)

    out = []
    for t, b, i in pieces:
        if out and out[-1][0].endswith(" ") and t.startswith(" "):
            t = t[1:]
        if t:
            out.append([t, b, i])
    if out:
        out[0][0] = out[0][0].lstrip()
        out[-1][0] = out[-1][0].rstrip()

    for t, b, i in out:
        if not t:
            continue
        r = p.add_run(t)
        r.bold = b or None
        r.italic = i or None
        if size:
            r.font.size = size
        if color:
            r.font.color.rgb = color


def heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt({1: 0, 2: 14, 3: 10}[level])
    p.paragraph_format.space_after = Pt({1: 10, 2: 4, 3: 2}[level])
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt({1: 20, 2: 13, 3: 11.5}[level])
    r.font.color.rgb = GREEN if level > 1 else INK
    return p


def label(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = GREEN
    return p


def add_table(doc, node):
    rows = [tr for tr in node.kids() if tr.tag == "tr"]
    if not rows:
        return
    ncols = max(
        sum(int(c.attrs.get("colspan", 1)) for c in tr.kids() if c.tag in ("td", "th"))
        for tr in rows
    )
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    table.autofit = True

    for tr in rows:
        cells = [c for c in tr.kids() if c.tag in ("td", "th")]
        row = table.add_row().cells
        header = bool(cells) and all(c.tag == "th" for c in cells)
        col = 0
        for c in cells:
            span = min(int(c.attrs.get("colspan", 1)), ncols - col)
            cell = row[col]
            if span > 1:
                cell = cell.merge(row[col + span - 1])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if c.attrs.get("style", "").find("right") >= 0:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_runs(
                p,
                c,
                bold=header,
                size=Pt(9.5),
                color=RGBColor(0xFF, 0xFF, 0xFF) if header else None,
            )
            if header:
                cell._tc.get_or_add_tcPr().append(
                    parse_xml(
                        r'<w:shd %s w:val="clear" w:color="auto" w:fill="0F8A5F"/>'
                        % nsdecls("w")
                    )
                )
            col += span
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_list(doc, node):
    ordered = node.tag == "ol"
    for li in node.kids():
        if li.tag != "li":
            continue
        sub = [c for c in li.kids() if c.tag == "p"]
        p = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
        p.paragraph_format.space_after = Pt(2)
        body = Node("span")
        body.children = [
            c for c in li.children if not (isinstance(c, Node) and c.tag == "p")
        ]
        add_runs(p, body, size=Pt(10.5))
        for s in sub:
            sp = doc.add_paragraph()
            sp.paragraph_format.left_indent = Mm(8)
            sp.paragraph_format.space_after = Pt(6)
            add_runs(sp, s, size=Pt(10), color=MUTED)


def block(node, doc, ctx):
    tag, cls = node.tag, node.cls

    if tag in ("script", "style"):
        return
    if "hr" in cls or "footer-note" in cls or "arrow-x" in cls:
        return

    if tag == "section":
        if not ctx["first"]:
            doc.add_page_break()
        ctx["first"] = False
        for c in node.kids():
            block(c, doc, ctx)
        return

    if "slide-head" in cls:
        for c in node.kids():
            if c.tag == "h2":
                heading(doc, c.text(), 1)
            elif c.tag == "span":
                label(doc, c.text())
        return

    if "title-logo" in cls:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(36)
        r = p.add_run(node.text())
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = GREEN
        return

    if tag == "h1":
        lines = [l.strip() for l in node.text().split("\n") if l.strip()]
        for i, line in enumerate(lines):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6 if i < len(lines) - 1 else 24)
            r = p.add_run(line)
            r.bold = True
            r.font.size = Pt(26 if i == 0 else 15)
            r.font.color.rgb = INK if i == 0 else GREEN
            if i:
                r.italic = True
        return

    if tag == "h2":
        heading(doc, node.text(), 1)
        return
    if tag == "h3":
        heading(doc, node.text(), 2)
        return

    if tag == "p":
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        add_runs(p, node, size=Pt(10.5))
        return

    if tag == "table":
        add_table(doc, node)
        return

    if tag in ("ul", "ol"):
        add_list(doc, node)
        return

    if "step" in cls:
        num = next((c.text() for c in node.kids() if "step-num" in c.cls), "")
        for c in node.kids():
            if c.tag == "h3":
                heading(doc, f"{num}. {c.text()}" if num else c.text(), 3)
            elif "days" in c.cls:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(10)
                r = p.add_run(c.text())
                r.bold = True
                r.font.size = Pt(9)
                r.font.color.rgb = AMBER
            else:
                block(c, doc, ctx)
        return

    if "big-num" in cls:
        main = "".join(c.text() for c in node.kids() if c.tag != "small").strip()
        small = "".join(c.text() for c in node.kids() if c.tag == "small").strip()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(main)
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = GREEN
        if small:
            sp = doc.add_paragraph()
            sp.paragraph_format.space_after = Pt(8)
            sr = sp.add_run(small)
            sr.font.size = Pt(9.5)
            sr.font.color.rgb = MUTED
        return

    if "node" in cls:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        b = next((c for c in node.kids() if c.tag == "b"), None)
        span = next((c for c in node.kids() if c.tag == "span"), None)
        if b:
            r = p.add_run(b.text())
            r.bold = True
            r.font.size = Pt(11)
        if span:
            p.add_run().add_break()
            r = p.add_run(span.text())
            r.font.size = Pt(9.5)
            r.font.color.rgb = MUTED
        return

    if "title-meta" in cls:
        for c in node.kids():
            b = next((x for x in c.kids() if x.tag == "b"), None)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            if b:
                p.add_run(f"{b.text()}: ").bold = True
            rest = c.text().replace(b.text(), "", 1).strip() if b else c.text()
            r = p.add_run(rest)
            r.font.size = Pt(10.5)
        return

    chips = [c for c in node.kids() if "chip" in c.cls]
    if chips:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(" · ".join(c.text() for c in chips))
        r.font.size = Pt(9.5)
        r.font.color.rgb = MUTED
        return

    for c in node.kids():
        block(c, doc, ctx)


def add_footer(section):
    p = section.footer.paragraphs[0]
    r = p.add_run("Гео Разум · Коммерческое предложение · стр. ")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    p._p.append(
        parse_xml(
            r'<w:fldSimple %s w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'
            % nsdecls("w")
        )
    )


def main(src="kp-slides.html", dst="kp.docx"):
    tree = Tree()
    with open(src, encoding="utf-8") as f:
        tree.feed(f.read())

    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.top_margin = Mm(20)
    s.bottom_margin = Mm(18)
    s.left_margin = Mm(18)
    s.right_margin = Mm(18)
    add_footer(s)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)

    html = next(n for n in tree.root.kids() if n.tag == "html")
    body = next(n for n in html.kids() if n.tag == "body")
    ctx = {"first": True}
    for section in body.kids():
        if section.tag == "section":
            block(section, doc, ctx)

    doc.save(dst)

    # check: содержимое на месте
    text = "\n".join(p.text for p in Document(dst).paragraphs)
    for expected in (
        "Коммерческое предложение",
        "330 000",
        "1 320 000",
        "275 ч",
        "50% — предоплата",
    ):
        assert expected in text, f"нет в документе: {expected}"
    print(
        f"{dst}: секций {len(Document(dst).sections)}, абзацев {len(text.splitlines())}"
    )


if __name__ == "__main__":
    main()
